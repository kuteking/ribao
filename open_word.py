from docx import Document
from tkinter import filedialog


def word_open():
    strtemp=''
    str_temp=''
    list_temp=[]
    filepath=filedialog.askopenfilename()
    doc=Document(filepath)
    tables=doc.tables
    table=tables[0]
    strtemp=table.cell(2,0).text
    for n in strtemp:
        if n !='\n':
            str_temp=str_temp+n
        else:
            if len(str_temp)>5:
                list_temp.append(str_temp)
            str_temp=''
    
    for nn in list_temp:
        if '累计' in nn:
            pass
        else:
            list_temp.remove(nn)
    bj = False
    str_temp=''
    data_end=[]
    for n in list_temp:
        for nn in n:
            if nn == '累':
                bj = True
            if '0'<=nn<='9' and bj == True:
                str_temp=str_temp+nn
        bj = False
        data_end.append(str_temp)
        str_temp=''
    return  filepath,tuple(data_end)
